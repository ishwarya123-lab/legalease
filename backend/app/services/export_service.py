import io
import re
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from docx import Document

def on_page_setup(canvas, doc):
    canvas.saveState()
    
    # Watermark Header
    canvas.setFont("Helvetica-Bold", 14)
    canvas.setFillColorRGB(0.8, 0.8, 0.8)
    canvas.drawCentredString(letter[0] / 2.0, letter[1] - 0.5 * inch, "CONFIDENTIAL")
    
    # Page Numbers
    canvas.setFont("Helvetica", 9)
    canvas.setFillColorRGB(0, 0, 0)
    page_num = canvas.getPageNumber()
    canvas.drawString(inch, 0.5 * inch, f"Page {page_num}")
    
    canvas.restoreState()

def generate_pdf(markdown_content: str, chart_bytes: bytes = None) -> io.BytesIO:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter,
                            rightMargin=inch, leftMargin=inch,
                            topMargin=inch, bottomMargin=inch)
    
    styles = getSampleStyleSheet()
    normal_style = styles['Normal']
    normal_style.fontSize = 11
    normal_style.leading = 14
    
    heading1_style = styles['Heading1']
    heading1_style.fontSize = 18
    heading1_style.leading = 22
    
    heading2_style = styles['Heading2']
    heading2_style.fontSize = 14
    heading2_style.leading = 18
    
    heading3_style = styles['Heading3']
    heading3_style.fontSize = 12
    heading3_style.leading = 16
    
    story = []
    
    lines = markdown_content.split('\n')
    for line in lines:
        if line.startswith('# '):
            story.append(Paragraph(line[2:], heading1_style))
            story.append(Spacer(1, 0.1 * inch))
        elif line.startswith('## '):
            story.append(Paragraph(line[3:], heading2_style))
            story.append(Spacer(1, 0.1 * inch))
        elif line.startswith('### '):
            story.append(Paragraph(line[4:], heading3_style))
            story.append(Spacer(1, 0.1 * inch))
        elif line.startswith('#### '):
            # Treat H4 as bold paragraph
            clean_text = line[5:]
            story.append(Paragraph(f"<b>{clean_text}</b>", normal_style))
            story.append(Spacer(1, 0.05 * inch))
        elif line.strip() == '' or line.strip().startswith('---') or line.strip() == '<br>':
            story.append(Spacer(1, 0.1 * inch))
        elif line.strip().startswith('|') and line.strip().endswith('|'):
            # It's a markdown table row, just clean it up a bit and render as text
            clean_row = line.strip().replace('|', '  ').replace('---', '').strip()
            if clean_row:
                story.append(Paragraph(clean_row, normal_style))
        else:
            import html
            line = html.escape(line)
            line = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', line)
            line = line.replace('₹', 'Rs. ')
            story.append(Paragraph(line, normal_style))
            
    if chart_bytes:
        story.append(Spacer(1, 0.2 * inch))
        chart_stream = io.BytesIO(chart_bytes)
        # Assuming typical size of 6x3 or similar, adjusting for page
        img = Image(chart_stream, width=5*inch, height=2.5*inch, kind='proportional')
        story.append(img)
            
    doc.build(story, onFirstPage=on_page_setup, onLaterPages=on_page_setup)
    buffer.seek(0)
    return buffer

def generate_docx(markdown_content: str, chart_bytes: bytes = None) -> io.BytesIO:
    doc = Document()
    
    lines = markdown_content.split('\n')
    for line in lines:
        if line.startswith('# '):
            doc.add_heading(line[2:], level=1)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        elif line.startswith('#### '):
            # Treat H4 as bold paragraph
            p = doc.add_paragraph()
            p.add_run(line[5:]).bold = True
        elif line.strip() == '' or line.strip().startswith('---') or line.strip() == '<br>':
            continue
        elif line.strip().startswith('|') and line.strip().endswith('|'):
            clean_row = line.strip().replace('|', '  ').replace('---', '').strip()
            if clean_row:
                doc.add_paragraph(clean_row)
        elif line.strip() != '':
            # Replace markdown bold for simple rendering
            clean_line = line.replace('**', '')
            doc.add_paragraph(clean_line)
            
    if chart_bytes:
        chart_stream = io.BytesIO(chart_bytes)
        doc.add_picture(chart_stream)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def generate_txt(markdown_content: str) -> io.BytesIO:
    buffer = io.BytesIO()
    buffer.write(markdown_content.encode('utf-8'))
    buffer.seek(0)
    return buffer
